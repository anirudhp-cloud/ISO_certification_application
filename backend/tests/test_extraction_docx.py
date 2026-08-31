# Regression tests for the docx reading-order/citation fix (M0 in
# ai_iso_plan_27_08_2026.md).
#
# The three bugs these pin down, all measured against a real uploaded policy
# before the fix:
#   1. tables were emitted after every paragraph, so a table a third of the way
#      down the document landed at the end — and because the flat text is rebuilt
#      from the chunks, that scrambled order was what the LLM read.
#   2. paragraph numbers counted only non-empty paragraphs, so citations drifted
#      from the real position (up to +33 on an 85-paragraph document).
#   3. heading context was discarded, leaving "Paragraph 51" as the only
#      locator — unverifiable, since Word shows no paragraph numbers.

from io import BytesIO

import pytest
from docx import Document as DocxDocument

from app.extraction.document_extraction import extract_text


def _build_docx() -> bytes:
    """A document shaped like the real ones: a control table before any heading,
    then headings with paragraphs, with a table sitting *between* two headings."""
    doc = DocxDocument()

    doc.add_paragraph("AI POLICY FOR ACME")
    control = doc.add_table(rows=2, cols=2)
    control.rows[0].cells[0].text = "Document ID"
    control.rows[0].cells[1].text = "AIMS-QMS-D-99"
    control.rows[1].cells[0].text = "Version"
    control.rows[1].cells[1].text = "2.0"

    doc.add_paragraph("")  # blank — must not consume a paragraph number
    doc.add_heading("1. Purpose", level=1)
    doc.add_paragraph("This policy sets out the commitment.")
    doc.add_paragraph("")  # blank
    doc.add_paragraph("It applies to all AI activities.")

    doc.add_heading("2. Roles", level=1)
    roles = doc.add_table(rows=2, cols=2)
    roles.rows[0].cells[0].text = "Role"
    roles.rows[0].cells[1].text = "Responsibility"
    roles.rows[1].cells[0].text = "CTO"
    roles.rows[1].cells[1].text = "Owns the AIMS"

    doc.add_heading("3. Review", level=1)
    doc.add_paragraph("Reviewed annually.")

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.fixture(scope="module")
def extracted() -> tuple[str, list[dict]]:
    _, text, chunks = extract_text("policy.docx", _build_docx())
    return text, chunks


def _locations(chunks: list[dict]) -> list[str]:
    return [c["location"] for c in chunks]


def test_table_stays_in_reading_order(extracted):
    """The roles table sits between headings 2 and 3 in the source, so it must be
    emitted there — not after every paragraph in the document."""
    _, chunks = extracted
    order = _locations(chunks)

    roles_row = next(i for i, loc in enumerate(order) if "table" in loc.lower() and "Roles" in loc)
    review_heading = next(i for i, loc in enumerate(order) if "3. Review" in loc)
    purpose_para = next(i for i, loc in enumerate(order) if "1. Purpose" in loc and "paragraph" in loc)

    assert purpose_para < roles_row < review_heading


def test_control_table_precedes_the_first_heading(extracted):
    """The document-control table is the 2nd body element; it must not be
    relocated to the end, which is where the old paragraphs-then-tables loop put it."""
    _, chunks = extracted
    order = _locations(chunks)

    control_row = next(i for i, loc in enumerate(order) if loc.startswith("Table"))
    first_heading = next(i for i, loc in enumerate(order) if loc.startswith("Heading"))

    assert control_row < first_heading


def test_flat_text_follows_document_order(extracted):
    """The flat text is what the LLM reads, so it must carry the same order."""
    text, _ = extracted
    assert text.index("AIMS-QMS-D-99") < text.index("This policy sets out")
    assert text.index("This policy sets out") < text.index("Owns the AIMS")
    assert text.index("Owns the AIMS") < text.index("Reviewed annually")


def test_paragraphs_are_cited_by_heading_and_restart_under_each(extracted):
    """Citations name a searchable heading, and the number is an offset within
    that heading rather than an absolute count that blank paragraphs skew."""
    _, chunks = extracted
    by_location = {c["location"]: c["text"] for c in chunks}

    assert by_location['Section "1. Purpose", paragraph 1'] == "This policy sets out the commitment."
    # The blank paragraph between them must not consume a number.
    assert by_location['Section "1. Purpose", paragraph 2'] == "It applies to all AI activities."
    # Numbering restarts under the next heading.
    assert by_location['Section "3. Review", paragraph 1'] == "Reviewed annually."


def test_blank_paragraphs_produce_no_chunks(extracted):
    _, chunks = extracted
    assert all(c["text"].strip() for c in chunks)


def test_table_rows_are_cited_with_table_and_row_number(extracted):
    _, chunks = extracted
    by_location = {c["location"]: c["text"] for c in chunks}
    assert by_location['Section "2. Roles", table 2, row 2'] == "CTO | Owns the AIMS"


def test_merged_cells_are_not_duplicated():
    """row.cells yields one entry per grid column, so a merged cell repeats.
    The row text must contain it once."""
    doc = DocxDocument()
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "SPANNED"
    table.rows[0].cells[0].merge(table.rows[0].cells[2])
    buffer = BytesIO()
    doc.save(buffer)

    _, _, chunks = extract_text("merged.docx", buffer.getvalue())
    row_texts = [c["text"] for c in chunks if "row" in c["location"].lower()]
    assert row_texts, "expected a table row chunk"
    assert row_texts[0].count("SPANNED") == 1


def test_header_text_is_extracted():
    """doc.paragraphs never sees headers, where controlled-document IDs live."""
    doc = DocxDocument()
    doc.sections[0].header.paragraphs[0].text = "AIMS-QMS-D-99 — Internal"
    doc.add_paragraph("Body text.")
    buffer = BytesIO()
    doc.save(buffer)

    _, text, chunks = extract_text("header.docx", buffer.getvalue())
    assert any(c["location"].startswith("Header") for c in chunks)
    assert "AIMS-QMS-D-99" in text


def test_quote_from_a_chunk_can_be_located():
    """End-to-end with the citation lookup the API depends on: a verbatim quote
    must resolve back to a heading-scoped location."""
    from app.ai.source_locator import locate_quote

    _, _, chunks = extract_text("policy.docx", _build_docx())
    assert locate_quote(chunks, "This policy sets out the commitment.") == 'Section "1. Purpose", paragraph 1'
    assert locate_quote(chunks, "Owns the AIMS") == 'Section "2. Roles", table 2, row 2'


# --- multi-chunk quote resolution ------------------------------------------
#
# Measured against a real policy document with real model output: 9 of 17 citations
# failed because the model quoted a heading plus the paragraphs beneath it, while
# chunks are one paragraph or table row each — so no single chunk contained the whole
# quote. Every one of those quotes WAS in the document. The locator now also matches
# chunk-inside-quote and cites where the passage starts, taking that document from
# 8/17 to 17/17.


def test_a_quote_spanning_several_chunks_resolves_to_where_it_starts():
    """What matters is that the citation lands in the right section, not which chunk
    inside it wins. A heading shorter than the length guard is skipped and the first
    paragraph is cited instead — equally findable, and arguably more useful, since
    that's where the substance is."""
    from app.ai.source_locator import locate_quote

    _, _, chunks = extract_text("policy.docx", _build_docx())
    spanning = "1. Purpose\nThis policy sets out the commitment.\nIt applies to all AI activities."

    location = locate_quote(chunks, spanning)
    assert location is not None
    assert "1. Purpose" in location


def test_a_long_heading_is_itself_cited_for_a_spanning_quote():
    """Above the length guard, the heading is the first matching chunk in document
    order — which is exactly where the quoted passage begins."""
    from app.ai.source_locator import locate_quote

    heading = "2.  Scope of the AI Management System"
    body = "The AIMS applies to all AI activities undertaken by the organisation."
    chunks = [
        {"location": f'Heading "{heading}"', "text": heading},
        {"location": f'Section "{heading}", paragraph 1', "text": body},
    ]

    assert locate_quote(chunks, f"{heading}\n{body}") == f'Heading "{heading}"'


def test_a_quote_spanning_paragraphs_without_the_heading_still_resolves():
    from app.ai.source_locator import locate_quote

    _, _, chunks = extract_text("policy.docx", _build_docx())
    spanning = "This policy sets out the commitment. It applies to all AI activities."

    assert locate_quote(chunks, spanning) == 'Section "1. Purpose", paragraph 1'


def test_a_short_incidental_line_inside_a_long_quote_is_not_used_as_the_location():
    """Guards the chunk-inside-quote pass: a stray short line that happens to appear
    within a long quote must not win over the real starting chunk."""
    from app.ai.source_locator import locate_quote

    chunks = [
        {"location": "Paragraph 1", "text": "Contents"},
        {"location": 'Section "1. Purpose", paragraph 1', "text": "This policy sets out the full commitment of the organisation."},
    ]
    quote = "This policy sets out the full commitment of the organisation. Contents follow."

    assert locate_quote(chunks, quote) == 'Section "1. Purpose", paragraph 1'


def test_a_quote_absent_from_the_document_still_resolves_to_nothing():
    from app.ai.source_locator import locate_quote

    _, _, chunks = extract_text("policy.docx", _build_docx())
    assert locate_quote(chunks, "a sentence that appears nowhere in this document at all") is None
