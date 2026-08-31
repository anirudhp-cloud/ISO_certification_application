# Format router — dispatches by file extension to the cheapest extraction path that
# works, per PHASE1_SCOPE.md section 3:
#   .txt            -> read directly
#   .docx / .pptx   -> local structured extraction (python-docx / python-pptx)
#   .xlsx           -> local structured extraction (openpyxl) — cell values, sheet by sheet
#   .pdf            -> native text-layer extraction first; OCR only if the text
#                      layer is empty/near-empty (i.e. it's a scan, not born-digital)
#
# Returns (extraction_method, text, chunks) where extraction_method is 'native' or
# 'ocr'. `chunks` is a list of {"location": str, "text": str} — the SAME text as
# `text`, just broken into position-tagged pieces (paragraph/slide/row/page) so a
# finding's quoted rationale can later be matched back to exactly where it came
# from (see app/ai/source_locator.py). `text` is still what's sent to the LLM,
# unchanged. Callers needing the OCR path when Document Intelligence isn't
# configured should catch DocumentIntelligenceNotConfigured (see
# document_intelligence_client.py).

import logging
from io import BytesIO

import fitz  # PyMuPDF
import openpyxl
from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation

from app.extraction.document_intelligence_client import ocr_document

logger = logging.getLogger(f"iso_platform.{__name__}")

# A PDF page with fewer than this many non-whitespace characters is treated as
# scanned/image-only rather than born-digital.
_MIN_NATIVE_CHARS_PER_PAGE = 20

# The formats extract_text() can actually read. Exposed so callers that ingest many
# files at once (the zip upload) can reject an unreadable entry up front and say so,
# rather than creating a document whose extraction silently yields nothing.
SUPPORTED_EXTENSIONS = frozenset({"txt", "docx", "pptx", "xlsx", "pdf"})


def extract_text(file_name: str, content: bytes) -> tuple[str, str, list[dict]]:
    extension = file_name.rsplit(".", 1)[-1].lower()
    if extension == "txt":
        text = content.decode("utf-8", errors="replace")
        chunks = [{"location": f"Line {i + 1}", "text": line} for i, line in enumerate(text.split("\n")) if line.strip()]
        logger.info("txt extraction: %d lines, %d chars", len(chunks), len(text))
        return "native", text, chunks
    if extension == "docx":
        return "native", *_extract_docx(content)
    if extension == "pptx":
        return "native", *_extract_pptx(content)
    if extension == "xlsx":
        return "native", *_extract_xlsx(content)
    if extension == "pdf":
        return _extract_pdf(content)
    raise ValueError(
        f"Unsupported file type for extraction: .{extension} "
        f"(supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))})"
    )


def _docx_row_text(row) -> str:
    """A table row's cells joined, with merged cells emitted once.

    python-docx's `row.cells` returns one entry per *grid column*, so a cell
    spanning three columns comes back three times and its text would repeat in
    the row string. De-duplicate on the underlying `w:tc` element, which is
    shared between the repeats.
    """
    seen: set[int] = set()
    cells = []
    for cell in row.cells:
        marker = id(cell._tc)
        if marker in seen:
            continue
        seen.add(marker)
        cells.append(cell.text.strip())
    return " | ".join(cells)


def _extract_docx(content: bytes) -> tuple[str, list[dict]]:
    """Walk the document body in true reading order, citing by heading.

    Two things this deliberately does *not* do, both of which were bugs:

      - It doesn't iterate `doc.paragraphs` then `doc.tables`. That emits every
        table after every paragraph, so a roles matrix sitting a third of the way
        down the document lands at the end — and since `text` below is rebuilt
        from the chunks, that scrambled order is what the LLM reads too. Walking
        `body.iterchildren()` and dispatching on the tag keeps tables where the
        author put them.

      - It doesn't cite an absolute paragraph number. Word displays no paragraph
        numbers and a .docx has no fixed pagination, so "Paragraph 51" is
        unverifiable even when the count is right (and the old count skipped
        blank paragraphs, so it drifted — up to +33 on a real 85-paragraph
        policy). A heading is searchable, so locations read
        `Section "5. Roles and Responsibilities", paragraph 3` and the paragraph
        number restarts under each heading.
    """
    doc = DocxDocument(BytesIO(content))
    chunks: list[dict] = []

    # Header/footer text carries controlled-document IDs, versions and
    # classification on documents like these — `doc.paragraphs` never sees it.
    for part_name, part in (("Header", "header"), ("Footer", "footer")):
        for section_num, section in enumerate(doc.sections, start=1):
            lines = [p.text.strip() for p in getattr(section, part).paragraphs if p.text.strip()]
            if lines:
                chunks.append({"location": f"{part_name} (section {section_num})", "text": " / ".join(lines)})

    heading = None  # most recent heading, or None until the first one
    para_num = 0  # resets under each heading
    table_num = 0

    def _where(detail: str) -> str:
        return f'Section "{heading}", {detail}' if heading else detail.capitalize()

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, doc)
            body_text = paragraph.text.strip()
            if not body_text:
                continue
            if paragraph.style is not None and paragraph.style.name.startswith(("Heading", "Title")):
                heading = body_text
                para_num = 0
                chunks.append({"location": f'Heading "{body_text}"', "text": body_text})
                continue
            para_num += 1
            chunks.append({"location": _where(f"paragraph {para_num}"), "text": body_text})
        elif child.tag == qn("w:tbl"):
            table_num += 1
            for row_num, row in enumerate(Table(child, doc).rows, start=1):
                row_text = _docx_row_text(row)
                if row_text.strip(" |"):
                    chunks.append({"location": _where(f"table {table_num}, row {row_num}"), "text": row_text})

    text = "\n".join(c["text"] for c in chunks)
    logger.info(
        "docx extraction: %d chunks (%d table(s)), %d chars", len(chunks), table_num, len(text)
    )
    return text, chunks


def _pptx_shape_text(shape) -> list[str]:
    """All text a shape carries, including the kinds a plain has_text_frame check misses.

    Two of those matter on real decks:
      - TABLES. A slide table has has_table, not has_text_frame, so the approval block
        on a controlled document ("Prepared By / Reviewed By / Approved By", with names,
        roles and dates) was extracted as nothing — and that is exactly the leadership
        and authority evidence clauses 5.1 and 5.3 ask for.
      - GROUPED shapes. A group has no text frame of its own; its children do, so the
        text sits one level down and needs recursing into.
    """
    if getattr(shape, "has_table", False):
        rows = []
        for row in shape.table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                rows.append(row_text)
        return rows

    if str(shape.shape_type).startswith("GROUP"):
        nested = []
        for child in shape.shapes:
            nested.extend(_pptx_shape_text(child))
        return nested

    if shape.has_text_frame and shape.text_frame.text.strip():
        return [shape.text_frame.text]

    return []


def _extract_pptx(content: bytes) -> tuple[str, list[dict]]:
    presentation = Presentation(BytesIO(content))
    chunks = []
    for slide_num, slide in enumerate(presentation.slides, start=1):
        # slide.shapes is z-order (back to front), which is the order shapes were
        # added, not the order a reader sees them. Sort top-to-bottom then
        # left-to-right so the flat text follows the slide as it's read. Shapes
        # with no position (top/left can be None) sort last rather than crashing.
        ordered = sorted(
            slide.shapes, key=lambda s: (s.top if s.top is not None else 1 << 31, s.left or 0)
        )
        for shape in ordered:
            for text in _pptx_shape_text(shape):
                chunks.append({"location": f"Slide {slide_num}", "text": text})
    text = "\n".join(c["text"] for c in chunks)
    logger.info("pptx extraction: %d slides, %d text chunks, %d chars", len(presentation.slides), len(chunks), len(text))
    return text, chunks


def _extract_xlsx(content: bytes) -> tuple[str, list[dict]]:
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=True, read_only=True)
    chunks = []
    for sheet in workbook.worksheets:
        for row_num, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = ["" if v is None else str(v) for v in row]
            row_text = " | ".join(cells)
            if any(cell.strip() for cell in cells):
                chunks.append({"location": f"Sheet '{sheet.title}', Row {row_num}", "text": row_text})
    # Rebuild the flat text with "Sheet: <name>" header lines for readability/LLM input.
    text_parts = []
    last_sheet = None
    for c in chunks:
        sheet_name = c["location"].split("'")[1]
        if sheet_name != last_sheet:
            text_parts.append(f"Sheet: {sheet_name}")
            last_sheet = sheet_name
        text_parts.append(c["text"])
    text = "\n".join(text_parts)
    logger.info("xlsx extraction: %d rows across %d sheet(s), %d chars", len(chunks), len(workbook.worksheets), len(text))
    return text, chunks


def _extract_pdf(content: bytes) -> tuple[str, str, list[dict]]:
    doc = fitz.open(stream=content, filetype="pdf")
    try:
        pages = [page.get_text() for page in doc]
        is_scanned = all(len(page.strip()) < _MIN_NATIVE_CHARS_PER_PAGE for page in pages)

        if not is_scanned:
            chunks = []
            for page_num, page in enumerate(doc, start=1):
                blocks = page.get_text("blocks")
                para_num = 0
                for block in blocks:
                    block_text = block[4].strip()
                    if not block_text:
                        continue
                    para_num += 1
                    chunks.append({"location": f"Page {page_num}, Paragraph {para_num}", "text": block_text})
            text = "\n".join(pages)
            logger.info("pdf extraction (native): %d pages, %d chunks, %d chars", len(pages), len(chunks), len(text))
            return "native", text, chunks
    finally:
        doc.close()

    # Scanned — whole-document OCR returns flat text only; no positional
    # structure available from this path today, so chunks stay empty (a
    # rationale quoted from an OCR'd document just won't get a source_location).
    logger.info("pdf extraction: scanned/image-only, falling back to OCR")
    return "ocr", ocr_document(content), []
